import json
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAVE_OPENPYXL = False

from finengine.verification import ManifestVerifier


def _supplement_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    inc = wb.active
    inc.title = "Income Statement"
    inc.append(["SAR (mn)", "FY 2023", "FY 2024", "YoY %", "1Q 2025", "FY 2025"])
    inc.append(["Net financing income", 24000, 27000, 0.125, 7000, 29000])
    inc.append(["Total operating income", 32000, 35000, 0.09, 9000, 39000])
    inc.append(["Operating expenses", -8000, -8500, 0.06, -2300, -9000])
    inc.append(["Provision for credit losses", -2000, -2100, 0.05, -600, -2300])
    inc.append(["Net income for the year", 19000, 21000, 0.10, 6000, 24000])
    inc.append(["Return on equity", 0.19, 0.20, 0.05, 0.19, 0.21])  # a ratio row - must be ignored
    bs = wb.create_sheet("Balance Sheet")
    bs.append(["SAR (mn)", "FY 2023", "FY 2024", "FY 2025"])
    bs.append(["Financing, net", 600000, 700000, 750000])
    bs.append(["Total assets", 900000, 1000000, 1040000])
    bs.append(["Total liabilities", 780000, 860000, 898000])
    bs.append(["Total equity", 120000, 140000, 142000])
    wb.save(path)


_MAPPING = {
    "company_id": "sa:9999",
    "source_url": "https://issuer.example/supplement.xlsx",
    "scale": "1000000",
    "sheets": {
        "Income Statement": {
            "Net financing income": ["net_financing_income", "fy"],
            "Total operating income": ["total_operating_income", "fy"],
            "Operating expenses": ["operating_expense_banking", "fy"],
            "Provision for credit losses": ["provision_expense", "fy"],
            "Net income for the year": ["net_income", "fy"],
        },
        "Balance Sheet": {
            "Financing, net": ["net_loans", "instant"],
            "Total assets": ["total_assets", "instant"],
            "Total liabilities": ["total_liabilities", "instant"],
            "Total equity": ["total_equity", "instant"],
        },
    },
}


@unittest.skipUnless(HAVE_OPENPYXL, "supplement reader needs the optional openpyxl extra")
class SupplementReaderTests(unittest.TestCase):
    def _read(self, directory: Path):
        from finengine.reading_xlsx import SupplementReader

        xlsx = directory / "supp.xlsx"
        _supplement_xlsx(xlsx)
        mapping = directory / "9999.json"
        mapping.write_text(json.dumps(_MAPPING), encoding="utf-8")
        return SupplementReader(xlsx, mapping).read(
            market="SA", symbol="9999", currency="SAR", filed_at="2026-02-04")

    def test_reads_multi_year_and_ignores_ratio_rows(self):
        with tempfile.TemporaryDirectory() as name:
            manifest = self._read(Path(name))
            years = sorted({f["fiscal_year"] for f in manifest["facts"]})
            self.assertEqual(years, [2023, 2024, 2025])          # FY only; 1Q skipped
            metrics = {f["metric"] for f in manifest["facts"]}
            self.assertIn("net_financing_income", metrics)
            self.assertIn("total_assets", metrics)
            self.assertNotIn("return_on_equity", metrics)        # issuer ratio row ignored

            by = {(f["metric"], f["fiscal_year"]): f for f in manifest["facts"]}
            fact = by[("total_assets", 2025)]
            self.assertEqual(fact["value"], "1040000")
            self.assertEqual(fact["scale"], "1000000")           # SAR (mn)
            self.assertEqual(fact["period_kind"], "instant")     # balance sheet -> instant
            self.assertEqual(by[("net_financing_income", 2025)]["period_kind"], "fy")

    def test_supplement_manifest_passes_verify(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            manifest = self._read(directory)
            imports = directory / "imports"
            imports.mkdir()
            (imports / "acme-supplement.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = ManifestVerifier(imports).verify()
            self.assertTrue(report["ok"], report["detail"])
            passed = {c["check"] for c in report["detail"] if c["status"] == "pass"}
            self.assertIn("balance_sheet: assets = liabilities + equity", passed)

    def test_flow_mapping_preserves_quarter_semantics(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            xlsx = directory / "supp.xlsx"
            _supplement_xlsx(xlsx)
            mapping = dict(_MAPPING)
            mapping["sheets"] = {
                "Income Statement": {
                    "Net income for the year": ["net_income", "flow"],
                }
            }
            mapping_path = directory / "9999.json"
            mapping_path.write_text(json.dumps(mapping), encoding="utf-8")
            from finengine.reading_xlsx import SupplementReader
            manifest = SupplementReader(xlsx, mapping_path).read(
                "SA", "9999", "SAR", "2026-02-04", period_kinds=("fy", "quarter")
            )
            q1 = next(f for f in manifest["facts"] if f["period_kind"] == "quarter")
            self.assertEqual(q1["period_start"], "2025-01-01")
            self.assertEqual(q1["period_end"], "2025-03-31")
            self.assertEqual(q1["fiscal_quarter"], 1)

    def test_flow_mapping_assigns_ytd_fiscal_quarter(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Income Statement"
            sheet.append(["SAR mn", "1H 2025", "9M 2025"])
            sheet.append(["Net income", 120, 190])
            xlsx = directory / "supp.xlsx"
            workbook.save(xlsx)
            workbook.close()
            mapping_path = directory / "9999.json"
            mapping_path.write_text(json.dumps({
                "scale": "1000000",
                "sheets": {"Income Statement": {
                    "Net income": ["net_income", "flow"],
                }},
            }), encoding="utf-8")
            from finengine.reading_xlsx import SupplementReader
            manifest = SupplementReader(xlsx, mapping_path).read(
                "SA", "9999", "SAR", "2025-10-30", period_kinds=("ytd",)
            )
            by_end = {fact["period_end"]: fact for fact in manifest["facts"]}
            self.assertEqual(by_end["2025-06-30"]["fiscal_quarter"], 2)
            self.assertEqual(by_end["2025-09-30"]["fiscal_quarter"], 3)

    def test_instant_mapping_keeps_interim_balance_dates(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Balance Sheet"
            sheet.append(["SAR mn", "FY 2024", "1Q 2025", "1H 2025"])
            sheet.append(["Total assets", 900, 940, 970])
            xlsx = directory / "supp.xlsx"
            workbook.save(xlsx)
            workbook.close()
            mapping_path = directory / "9999.json"
            mapping_path.write_text(json.dumps({
                "sheets": {"Balance Sheet": {
                    "Total assets": ["total_assets", "instant"],
                }},
            }), encoding="utf-8")
            from finengine.reading_xlsx import SupplementReader
            manifest = SupplementReader(xlsx, mapping_path).read(
                "SA", "9999", "SAR", "2025-08-01",
                period_kinds=("fy", "quarter", "ytd"),
            )
            self.assertEqual(
                [(fact["period_end"], fact["period_kind"]) for fact in manifest["facts"]],
                [("2024-12-31", "instant"), ("2025-03-31", "instant"),
                 ("2025-06-30", "instant")],
            )

    def test_row_options_override_units_and_retain_excluded_source_values(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Income Statement"
            sheet.append(["SAR mn", "FY 2024", "FY 2025"])
            sheet.append(["Earnings per share", 4.5, 5.5])
            xlsx = directory / "supp.xlsx"
            workbook.save(xlsx); workbook.close()
            mapping_path = directory / "9999.json"
            mapping_path.write_text(json.dumps({
                "scale": "1000000",
                "sheets": {"Income Statement": {
                    "Earnings per share": ["eps_diluted", "flow", {
                        "scale": 1, "unit": "SAR/share",
                        "exclude_period_ends": ["2025-12-31"],
                        "exclude_reason": "audited filing takes precedence",
                    }],
                }},
            }), encoding="utf-8")
            from finengine.reading_xlsx import SupplementReader
            manifest = SupplementReader(xlsx, mapping_path).read(
                "SA", "9999", "SAR", "2026-02-04"
            )
            self.assertEqual(manifest["facts"][0]["scale"], "1")
            self.assertEqual(manifest["facts"][0]["unit"], "SAR/share")
            self.assertEqual(manifest["excluded_facts"][0]["period_end"], "2025-12-31")
            self.assertEqual(manifest["excluded_facts"][0]["value"], "5.5")
            self.assertEqual(manifest["excluded_facts"][0]["reason"],
                             "audited filing takes precedence")


if __name__ == "__main__":
    unittest.main()

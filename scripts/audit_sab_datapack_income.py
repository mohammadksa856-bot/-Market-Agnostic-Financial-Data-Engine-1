"""Independent cell-level reconciliation; no production writes or workbook edits."""
import json
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

root = Path('/app/state/reports/sab-period-correction-review/datapack-review')
output = root / 'income-cell-reconciliation-v1.json'
assert not output.exists(), 'Preserve audit checkpoint'
digests = ['76acffa163e24985e1c3b9ea9a2d986656291d3698da73b3608df030956c5d0c',
           '300a90f7b1ead3c808e7533087326f82bb137678422d8e8e438720f627f8aee7']
results = []
for digest in digests:
    path = Path('/app/state/raw/SA/1060/documents') / (digest + '.xlsx')
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook['Income Statement& Balance Sheet']
    cells = list(sheet.iter_rows(values_only=True))
    for col in range(2, 19):
        period = cells[3][col].date().isoformat()
        def amount(row):
            return Decimal(str(cells[row-1][col])).quantize(Decimal('0.001'))
        equations = {
            'net_commission': (8, [6, 7]),
            'operating_income': (12, [8, 9, 10, 11]),
            'operating_expenses': (17, [13, 14, 15, 16]),
            'before_provisions': (18, [12, 17]),
            'operating_profit': (20, [18, 19]),
            'before_tax': (22, [20, 21]),
            'continuing_income': (24, [22, 23]),
            'total_income_with_discontinued': (26, [24, 25]),
            'income_attribution': (30, [28, 29]),
            'after_coupon': (32, [30, 31]),
        }
        checks = []
        for name, (total, parts) in equations.items():
            if any(not isinstance(cells[row-1][col], (int, float)) for row in [total, *parts]):
                checks.append({'equation': name, 'status': 'missing_or_non_numeric_source_cell',
                               'total_row': total, 'component_rows': parts, 'within_printed_precision': False})
                continue
            delta = amount(total) - sum((amount(row) for row in parts), Decimal(0))
            checks.append({'equation': name, 'total_row': total, 'component_rows': parts,
                           'delta_SAR': str(delta * 1000000),
                           'within_printed_precision': abs(delta) <= Decimal('0.002')})
        results.append({'hash': digest, 'period_end': period, 'checks': checks,
                        'duplicate_net_income_rows': [26, 30],
                        'duplicate_net_income_delta_SAR': str((amount(26)-amount(30))*1000000)
                        if all(isinstance(cells[r-1][col], (int,float)) for r in [26,30]) else None})
    workbook.close()
report = {'production_modified': False, 'workbooks': len(digests), 'columns': len(results),
          'checks': sum(len(r['checks']) for r in results),
          'failed_equations': sum(not c['within_printed_precision'] for r in results for c in r['checks']),
          'different_duplicate_income_columns': sum(r['duplicate_net_income_delta_SAR'] is not None and Decimal(r['duplicate_net_income_delta_SAR']) != 0 for r in results),
          'results': results, 'publication_approved': False,
          'next_gate': 'Resolve attribution discrepancies against audited statements; do not map repeated labels globally.'}
output.write_text(json.dumps(report, indent=2))
print(json.dumps({k:v for k,v in report.items() if k != 'results'}))

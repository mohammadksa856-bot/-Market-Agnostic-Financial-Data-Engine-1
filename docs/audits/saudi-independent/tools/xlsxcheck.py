"""Independent check of data-supplement manifests against the archived XLSX workbook (read-only)."""
import sys, json, os, re, collections, difflib, warnings
warnings.filterwarnings("ignore")
import openpyxl

def norm(s):
    s = re.sub(r"[^a-z0-9 ]+", " ", str(s).lower().replace("�", "'"))
    return re.sub(r"\s+", " ", s).strip()

QEND = {1: "03-31", 2: "06-30", 3: "09-30", 4: "12-31"}

def parse_header(h):
    if h is None:
        return None
    t = str(h).strip().upper().replace("\n", " ")
    m = re.fullmatch(r"FY\s*'?(\d{4})", t)
    if m: return ("fy", f"{m.group(1)}-12-31")
    m = re.fullmatch(r"([1-4])Q\s*'?(\d{4})", t)
    if m: return ("quarter", f"{m.group(2)}-{QEND[int(m.group(1))]}")
    m = re.fullmatch(r"(1H|6M|H1)\s*'?(\d{4})", t)
    if m: return ("ytd", f"{m.group(2)}-06-30")
    m = re.fullmatch(r"(9M|3Q YTD)\s*'?(\d{4})", t)
    if m: return ("ytd", f"{m.group(2)}-09-30")
    return None

def load_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = None
    for i, r in enumerate(rows[:12]):
        if r and r[0] and "sar" in str(r[0]).lower() and sum(1 for c in r[1:] if parse_header(c)) >= 2:
            hdr_i = i; break
    if hdr_i is None:
        return None
    cols = {}
    for j, c in enumerate(rows[hdr_i]):
        p = parse_header(c)
        if p: cols.setdefault(p, []).append(j)
    data = {}
    for k, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        if r and r[0] is not None:
            data.setdefault(norm(r[0]), []).append((k, r))
    return {"cols": cols, "data": data, "hdr_row": hdr_i + 1}

def check(manifest_path, xlsx_path):
    x = json.load(open(manifest_path, encoding="utf8"))
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = {}
    for ws in wb:
        if ws.sheet_state != "visible" or re.search(r"segment|disclaimer|content|operating", ws.title, re.I):
            continue
        if ws.max_row > 400 or ws.max_column > 120:
            continue
        s = load_sheet(ws)
        if s: sheets[ws.title] = s
    out = []
    for kind, lst in (("published", x["facts"]), ("excluded", x.get("excluded_facts", []))):
        for f in lst:
            r = {"kind": kind, "metric": f["metric"], "label": f["source_label"], "period_end": f["period_end"], "period_kind": f["period_kind"], "value": f["value"], "reason": (f.get("reason") or "")[:60]}
            want_kind = "fy" if f["period_kind"] in ("fy",) else f["period_kind"]
            keys = [("quarter" if f["period_kind"] == "instant" else f["period_kind"], f["period_end"])]
            if f["period_kind"] == "instant":
                keys = [("quarter", f["period_end"]), ("fy", f["period_end"]), ("ytd", f["period_end"])]
            cands = []
            for st, s_ in sheets.items():
                rows = s_["data"].get(norm(f["source_label"]))
                if not rows:
                    cand = difflib.get_close_matches(norm(f["source_label"]), list(s_["data"]), n=1, cutoff=0.8)
                    rows = s_["data"][cand[0]] if cand else None
                if not rows: continue
                for (rn, row) in rows:
                    for kk in keys:
                        for j in s_["cols"].get(kk, []):
                            if j < len(row) and isinstance(row[j], (int, float)):
                                cands.append((st, rn, j, row[j], kk))
            if not cands:
                r["status"] = "NOT_FOUND"; out.append(r); continue
            fv0 = float(f["value"])
            good = [c for c in cands if abs(abs(fv0) - abs(float(c[3]))) <= max(5e-4, abs(float(c[3])) * 1e-6)]
            found = good[0] if good else cands[0]
            r["n_candidates"] = len(cands); r["all_candidate_values"] = sorted({round(float(c[3]), 3) for c in cands})
            st, rn, j, cell, kk = found
            fv = float(f["value"]); scale = float(f.get("scale") or 1)
            cv = float(cell)
            ok = abs(abs(fv) - abs(cv)) <= max(5e-4, abs(cv) * 1e-6)
            r.update({"sheet": st, "row": rn, "col": openpyxl.utils.get_column_letter(j + 1), "cell_value": cell, "status": "OK" if ok else "VALUE_MISMATCH", "sign_same": (fv >= 0) == (cv >= 0) or fv == 0})
            out.append(r)
    return out

if __name__ == "__main__":
    a = json.load(open("data/raw/archive-index.json", encoding="utf8"))["artifacts"]
    outdir = sys.argv[1]; os.makedirs(outdir, exist_ok=True)
    for mf in sys.argv[2:]:
        b = os.path.basename(mf)
        paths = [z["local_path"] for z in a if b in z["metadata"].get("manifests", [])]
        res = check(mf, paths[0])
        json.dump(res, open(os.path.join(outdir, b), "w", encoding="utf8"), ensure_ascii=False)
        print(b, dict(collections.Counter((z["status"], z.get("sign_same")) for z in res)))
